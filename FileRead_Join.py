from pathlib import Path
import tkinter as tk
from tkinter import filedialog
from openpyxl import load_workbook
from datetime import datetime


def get_file_list(folder):
    files = [f.name for f in folder.iterdir() if f.is_file() and not f.name.startswith('.')]
    files.sort(key=str.lower)
    return files


def choose_excel_file():
    root = tk.Tk()
    root.withdraw()

    return filedialog.askopenfilename(
        title="Select Excel file",
        filetypes=[("Excel files", "*.xlsx")]
    )


def choose_sheet_name(excel_path):
    wb = load_workbook(excel_path)
    sheet_names = wb.sheetnames

    print("\nAvailable sheets:")
    for i, name in enumerate(sheet_names, start=1):
        print(f"{i}. {name}")

    while True:
        choice = input("\nEnter the sheet number or sheet name: ").strip()

        if choice.isdigit():
            index = int(choice) - 1
            if 0 <= index < len(sheet_names):
                return sheet_names[index]

        if choice in sheet_names:
            return choice

        print("Invalid choice. Please try again.")


def choose_slide_range():
    while True:
        try:
            start_pos = int(input("Enter starting slide position in box: ").strip())
            end_pos = int(input("Enter ending slide position in box: ").strip())

            if start_pos < 1 or end_pos < start_pos:
                print("Invalid range. Try again.")
                continue

            return start_pos, end_pos

        except ValueError:
            print("Please enter whole numbers only.")


def write_filenames_to_sheet(excel_path, sheet_name, files, start_pos, end_pos, col="J"):
    from datetime import datetime
    from openpyxl import load_workbook

    wb = load_workbook(excel_path)
    ws = wb[sheet_name]

    available_slots = end_pos - start_pos + 1
    files_to_write = files[:available_slots]

    for i, filename in enumerate(files_to_write):
        slide_position = start_pos + i
        excel_row = slide_position + 3
        ws[f"{col}{excel_row}"] = filename

    # 👇 timestamped filename
    timestamp = datetime.now().strftime("%Y%m%d_%H%M")
    output_path = excel_path.replace(".xlsx", f"_updated_{timestamp}.xlsx")

    wb.save(output_path)

    return output_path, len(files_to_write)


def main():
    folder = Path("/Users/bb5117/Library/CloudStorage/OneDrive-ImperialCollegeLondon/LabWork/Microscope Images/260416")
    col = "J"

    files = get_file_list(folder)

    excel_path = choose_excel_file()
    if not excel_path:
        print("No file selected.")
        return

    print("Selected file:", excel_path)

    sheet_name = choose_sheet_name(excel_path)
    print("Selected sheet:", sheet_name)

    start_pos, end_pos = choose_slide_range()
    print(f"Writing files for slide positions {start_pos} to {end_pos}")

    output_path, n_written = write_filenames_to_sheet(
        excel_path=excel_path,
        sheet_name=sheet_name,
        files=files,
        start_pos=start_pos,
        end_pos=end_pos,
        col=col
    )

    print(f"Wrote {n_written} filenames to column {col}")
    print(f"Saved updated file: {output_path}")


if __name__ == "__main__":
    main()